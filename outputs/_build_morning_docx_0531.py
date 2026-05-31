# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2C, 0x5A, 0x8C)
RED  = RGBColor(0xC0, 0x39, 0x2B)
GREEN= RGBColor(0x1E, 0x7A, 0x46)
GREY = RGBColor(0x88, 0x88, 0x88)

def set_font(run, size=10.5, bold=False, color=None, italic=False):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_title(text, sub=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), 18, True, NAVY)
    if sub:
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p2.add_run(sub), 10, False, GREY)

def add_h2(text):
    p = doc.add_paragraph(); set_font(p.add_run(text), 14, True, NAVY)
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)

def add_h3(text):
    p = doc.add_paragraph(); set_font(p.add_run(text), 11.5, True, BLUE)
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)

def add_para(segs, space_before=0):
    p = doc.add_paragraph()
    if space_before: p.paragraph_format.space_before = Pt(space_before)
    for seg in segs:
        text, bold = seg[0], seg[1]
        color = seg[2] if len(seg) > 2 else None
        set_font(p.add_run(text), 10.5, bold, color)
    return p

def add_bullet(segs):
    p = doc.add_paragraph(style='List Bullet')
    for seg in segs:
        text, bold = seg[0], seg[1]
        color = seg[2] if len(seg) > 2 else None
        set_font(p.add_run(text), 10.5, bold, color)

def add_quote(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
    set_font(p.add_run(text), 9.5, False, RGBColor(0x66,0x66,0x66), italic=True)

def add_hr():
    p = doc.add_paragraph(); pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pBdr.makeelement(qn('w:bottom'), {qn('w:val'):'single', qn('w:sz'):'6', qn('w:space'):'1', qn('w:color'):'BFBFBF'})
    pBdr.append(bottom); pPr.append(pBdr)

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    for i, t in enumerate(headers):
        hdr[i].text = ''; set_font(hdr[i].paragraphs[0].add_run(t), 10.5, True, RGBColor(0xFF,0xFF,0xFF))
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''; set_font(cells[i].paragraphs[0].add_run(str(val)), 10)

# ===== 标题 =====
add_title('交易系统早间日报（链 A 产出）',
          '2026-05-31  11:00 评分点  |  Trading-System v2026.5  |  output_agent')
add_quote('数据采集 5/5 ✅：BlockBeats（Chrome MCP 实时抓取 09:32）/ Exocharts（截图 BTCUSD 日线 footprint）/ Coinglass（截图 BTCUSDT 1h）/ CME FedWatch（截图 6.17 FOMC）/ TradingView 15m（截图，盘面执行最终 Check）')
add_quote('持仓校准：直读 Notion REST「(主)/(Maker)交易Check」，非沿用本地 state。')
add_hr()

# ============================================================
#  用户 Notion 格式正文（可直接复制粘贴）
# ============================================================
add_h2('【主账户】2026-05-31  （Notion 格式 · 可直接复制）')

add_para([('不能做：', True), ('做多', True, RED), ('（已持多仓，不在区间中部追多）', False)])

add_para([('基本面：', True), ('利多', True, GREEN),
          ('（Binance 6/1 新增美股交易叙事引爆 BNB/BSC 生态全面反弹 + Loracle 等空头回补平 BTC 空单）', False),
          ('  价格部分兑现', True),
          ('（BTC 仅 +0.35%，涨幅集中 BNB/HYPE）', False)], space_before=4)

add_para([('技术面：', True)], space_before=4)
add_bullet([('盘口（合约）：', True), ('CVD 横盘　价格升　资金费率横盘降温（0.0033 正值）　持仓降后企稳', False)])
add_bullet([('解读：', True), ('中性偏多', True, GREEN),
            ('（去杠杆出清后低位企稳回升；但现货买盘弱、空头回补性质，价升 CVD 未确认 → 存隐性背离）', False)])
add_bullet([('盘面 · 日线级别：', True), ('POC 连续 4 日抬升（72,882→74,046）+ Delta 连续 2 日翻正 + 卖压衰竭，下行动能枯竭，筑底向上修复　', False), ('中性偏多', True, GREEN)])
add_bullet([('盘面 · 小时级别：', True), ('1h 反弹企稳站上 SMA20，价升但 CVD 未同步强升、买盘力度待确认　', False), ('中性偏多', True, GREEN)])

add_para([('趋势确认：', True), ('震荡结构（短线，空头下行动能枯竭转筑底修复）', True, BLUE)], space_before=4)
add_para([('日内支撑：', True), ('73,000-73,600', False), ('　　　日内压力：', True), ('74,400-74,750（次级 76,000 多空分界）', False)])
add_para([('策略：', True),
          ('中线计划保持不变 — 多仓 @73,500 持仓中，浮盈 +4.6%（仓位口径），未到 1.5R 防守层未激活，止损 71,000，', False),
          ('不新增', True, RED),
          ('；待有效突破 76,000 多空分界回踩 或 回踩 73k 关键区获支撑 再议加仓。', False)])

add_hr()
add_h2('【Maker】2026-05-31  （Notion 格式 · 可直接复制）')

add_para([('BTC（策略）：', True), ('观望', True, RED), ('（总评分 67 未达 70 开仓门槛）', False)])
add_bullet([('执行概率 0%', True)])
add_bullet([('反弹多触发（二选一 + CVD 转强确认）：', True), ('① 回踩 73,000-73,600 共振支撑获撑；② 放量突破 74,400 箱体上沿 → 76,000 后回踩确认', False)])
add_bullet([('当前价 74,082 逼近箱体上沿压力 74,400（+0.4%），', False), ('追多 = 买进阻力，入场位置不优，不在边缘不开', True, RED), ('（宪法：入场位置要求极高）', False)])
add_bullet([('不追空：', True), ('今日筑底回升 + 基本面利多 + 空头回补，且 74,400 被拒空与利多冲突，底部追空风险最大', False)])

add_hr()

# ============================================================
#  系统评分明细
# ============================================================
add_h2('一、11:00 评分点')
add_table(['维度', '满分', '实得', '备注'],
          [['盘口', '70', '50', '非极端组合：去杠杆出清后低位企稳回升（较 5/29-30 方向转明确 +5，CVD 背离压制）'],
           ['盘面', '30', '17', '详见下表（含 15m 最终盘面执行 Check）'],
           ['总评分', '100', '67', '< 70 门槛 → Maker 观望（5/29=63 / 5/30=62 / 5/31=67 逐步抬升）']])

add_h3('盘面 30 细分')
add_table(['项', '满分', '实得', '理由'],
          [['支撑/压力接近度', '10', '5', '15m：价 74,082 逼近箱体上沿压力 74,400（+0.4%）；但对做多方向是买进阻力，非优质多头入场位'],
           ['多空流动性掠夺', '10', '5', '72k 下方止损 5/28 已扫，向上修复未掠夺上方流动性'],
           ['SMA20 突破/回踩', '5', '3', '4H 站上 SMA20（积极），但日线仍空头未修复（保守）'],
           ['量价关系', '3', '2', '日波幅极窄 + 5/30 量能 59M 萎缩，无放量突破'],
           ['斐波/趋势线接近度', '2', '2', '贴大级别上升通道下轨反弹（结构性位置）']])

add_hr()

# ============================================================
#  五源数据支撑
# ============================================================
add_h2('二、五源数据支撑')

add_h3('① BlockBeats 基本面（实时抓取，截至 09:32）')
for ts, content, tag, color in [
    ('09:30', 'BNB/HYPE 引领加密全面反弹，BTC 再破 7.4 万，BSC 生态拉升（Binance 将新增股票交易引爆）', '利多·核心催化', GREEN),
    ('09:18', 'Loracle「空转多」平仓 BTC/HYPE/TON 等空单，开 ZEC/ASTER/TON 多单', '利多·空头回补', GREEN),
    ('08:53/01:07', 'Binance 明日（6/1）或新增真实美股交易 + bStocks 链上证券代币化（Alpaca 托管）', '利多·生态叙事', GREEN),
    ('09:32/08:40', '特朗普收紧美伊协议条款（核材料+霍尔木兹），60 天窗口未敲定，伊朗约 3 天回应', '利空·地缘未定', RED),
    ('其他', '柬埔寨绑架案 / 欧冠 Polymarket×2 / Loracle 减仓 HYPE', '中性', GREY),
]:
    p = doc.add_paragraph(style='List Bullet')
    set_font(p.add_run('['+ts+'] '), 10, True)
    set_font(p.add_run(content+'  '), 10)
    set_font(p.add_run('→ '+tag), 10, True, color)
add_para([('聚合：', True), ('利多 4 / 利空 2（同一地缘事件）/ 中性 4 → 差 ≥ 2 → ', False), ('利多', True, GREEN),
          ('。核心 = Binance 美股叙事引爆生态全面反弹 + 空头回补；制约 = 美伊协议美方强硬未定；价格部分兑现。', False)], space_before=4)

add_h3('② Exocharts 盘面（BTCUSD 日线 footprint）')
add_table(['日期', 'POC', 'Delta', 'CVD', '日波幅'],
          [['5/27', '75,318', '-6.3M', '3.5M', '1,934.9'],
           ['5/28', '72,882', '-19M', '-15M', '2,010.9（破位）'],
           ['5/29', '73,494', '-16M', '-31M', '1,832.8'],
           ['5/30', '73,686', '+23M', '-8.5M', '921.4'],
           ['5/31', '74,046', '+5.6M', '-3M', '417.9（今日刚开盘）']])
add_para([('POC 连续 4 日阶梯抬升 + Delta 连续 2 日翻正 + 日波幅收窄 = ', False), ('卖压衰竭、买方回归、低位筑底向上修复', True, GREEN), ('（较 5/30 横盘筑底再进一步）。', False)], space_before=4)

add_h3('③ Coinglass 盘口（BTCUSDT 1h）')
add_table(['维度', '读数', '判定'],
          [['CVD（累积成交量增量）', '-14.283K（深跌后横盘微抬）', '横盘/微升'],
           ['价格', '74,156（+0.22%，最新拉升）', '升'],
           ['资金费率（持仓加权）', '0.0033（从 ~0.010 大幅降温）', '正值偏冷'],
           ['持仓 OI', '54.765K（从 62K 回落企稳）', '降后企稳']])
add_para([('16 组合 = 去杠杆出清后低位企稳回升（中性偏多，非极端）。', True),
          ('⚠️ 价升但 CVD 未同步强升 = 现货买盘弱、偏空头回补性质，可持续性待确认。', False, RED)], space_before=4)

add_h3('④ CME FedWatch（6/17 FOMC）')
add_table(['目标利率', 'NOW', '1日', '1周', '1月'],
          [['325-350 降息', '0.4%', '0.4%', '0.0%', '5.6%'],
           ['350-375 不变', '99.6%', '99.6%', '96.0%', '94.4%'],
           ['375-400 加息', '0.0%', '0.0%', '4.0%', '0.0%']])
add_para([('维持不变几成定局；周比加息 4%→0%（去鹰）、月比降息 5.6%→0.4%（去鸽）→ ', False), ('macro 净中性（status quo），同 5/30。', True)], space_before=4)

add_h3('⑤ TradingView 15m（盘面执行 Check · 最终/最重要）')
add_para([('K 线：', True), ('开 74,058 / 高 74,082 / 低 74,058 / 收 74,081.98（09:51，+0.03%）', False)])
add_table(['类型', '价位', '说明'],
          [['空头流动性（信号2）', '~78,800', '更上方目标'],
           ['强压/多空分界', '76,000', '箱体上方开阔区终点'],
           ['压力 5.2(9) · 箱体上沿', '74,400', '当前最近阻力（价在其下方 +0.4%）'],
           ['当前价', '74,082', '已突破箱体中部虚线 73,600，逼近上沿压力'],
           ['箱体下沿 / 支撑 5.2', '72,800-73,200', '5/28 起本地箱体底（虚线 73,200/72,800）'],
           ['支撑 5.2(7)(9)', '72,000-72,400', '破位则趋势空延续']])
add_para([('盘面执行结论：', True),
          ('15m 看价格已顶到本地箱体（72,800-74,400）上沿、逼近 74,400 压力（+0.4%）。', False),
          ('对做多方向 = 买进阻力，入场位置不优 → 观望。', True, RED)], space_before=4)
add_bullet([('高质量入场触发（二选一）：', True), ('① 回踩 73,000-73,600 共振支撑获撑（反弹多）；② 放量突破 74,400 → 76,000 后回踩确认', False)])
add_bullet([('反弹空（74,400 被拒）与「利多 + 不追空」冲突，不取', False)])

add_hr()

# ============================================================
#  持仓校准
# ============================================================
add_h2('三、持仓校准（直读 Notion REST）')
add_bullet([('主账户：', True), ('5/28 做多 @ 73,500（底仓 30% / $1000 / 5x / 盈亏比 1:3），状态【持仓中】。TP1 78,000 减 50%，止损有效跌破 71,000。', False)])
add_bullet([('　浮盈防守：', True), ('① 75,500 设日线低点止损+取消加仓　② 76,200 设保本损', False)])
add_bullet([('　当前浮盈：', True), ('现价 ~74,180 → +0.93% × 5x ≈ ', False), ('+4.6%（仓位口径）', True, GREEN), ('；1.5R 触发位 = 77,250，远未到 → 防守层未激活', False)])
add_bullet([('Maker：', True), ('5/29 做多已平仓（-$9/2%）→ ', False), ('当前无持仓', True)])

add_hr()

# ============================================================
#  交易宪法对照
# ============================================================
add_h2('四、交易宪法对照')
for item in [
    '非系统不做 / 非 Maker 不开 / 违规即封禁交易资格',
    '浮盈首次 ≥ +1.5R 才激活防守层（主账户当前 +4.6%，未到）',
    '亏损砍短 盈利放满；单笔止损 3% 以内',
    '入场位置要求极高 → 当前 74,082 逼近 74,400 压力（追多买阻力）不满足 → 观望',
]:
    add_bullet([(item, False)])

# Footer
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run('Generated by output_agent (chain A)  |  2026-05-31 11:00  |  Trading-System v2026.5'), 9, False, RGBColor(0x99,0x99,0x99), italic=True)

out_path = r'C:\Users\31702\Desktop\Trading-System\outputs\2026-05-31_早间日报.docx'
doc.save(out_path)
print('SAVED:', out_path)
