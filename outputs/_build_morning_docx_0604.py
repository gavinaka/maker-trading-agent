# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Microsoft YaHei'
st.font.size = Pt(10.5)
st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def setfont(run, size=10.5, bold=False, color=None):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

def H(text, size=13):
    p = doc.add_paragraph()
    r = p.add_run(text)
    setfont(r, size=size, bold=True, color=RGBColor(0x1F,0x49,0x7D))
    p.space_after = Pt(4)
    return p

def line(text, bold=False, color=None, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    setfont(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(2)
    return p

RED = RGBColor(0xC0,0x00,0x00)
ORANGE = RGBColor(0xB0,0x60,0x00)
GREEN = RGBColor(0x1E,0x7A,0x32)

p = doc.add_paragraph()
r = p.add_run('2026-06-04 早间日报（系统独立稿）')
setfont(r, size=16, bold=True, color=RGBColor(0x00,0x00,0x00))
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

line('门禁 6/4＝可交易（今日行仅填情绪平静度＝平静，无「不可交易/收缩机制」标记 → 按可交易处理）；心力最新 6/3＝专注7/消耗3/无异常。总评分 71（盘口50+盘面21）≥70 → 必挂 Maker 单于高质量边缘。现价 64,114（日线收）。', color=RGBColor(0x55,0x55,0x55), size=9)

# ===== 正文 =====
H('① 日报正文')
line('不能做：（追）做空', bold=True, color=RED)
p = doc.add_paragraph()
r1 = p.add_run('基本面：'); setfont(r1, bold=True)
r2 = p.add_run('中性偏空'); setfont(r2, bold=True, color=ORANGE)
r3 = p.add_run('（周四开门黑、BTC跌破6.5万跌势延续，过去24h全网爆仓9.84亿、多单爆仓8.04亿=多头继续被血洗；美股三大股指收跌、微软英伟达跌逾3%=risk-off传导，日韩股市低开；美股资金分化270亿流入科技、其余板块流出40亿=币市持续失血。对冲项=地缘转缓和：特朗普称对伊谈判"非常顺利"协议本周末或达成、伊朗四阶段计划首阶段全面停火，但今晨仍称打击美军指挥舰、缓和未落实）'); setfont(r3)
r4 = p.add_run('价格部分兑现'); setfont(r4, bold=True, color=GREEN)

# 基本面配图（BlockBeats 快讯卡）
for _img in ['2026-06-04_1_btc.png','2026-06-04_2_equity.png','2026-06-04_3_iran.png']:
    _p = doc.add_paragraph(); _p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _p.add_run().add_picture(r'C:\Users\31702\Desktop\Trading-System-Code\outputs\blockbeats\\'+_img, width=Inches(5.2))

line('技术面：', bold=True)
line('盘口（合约）：', bold=True)
line('CVD下降　价格下降　资金费率下降（为正值）　持仓横盘')
p = doc.add_paragraph()
ra = p.add_run('解读：'); setfont(ra, bold=True)
rb = p.add_run('利空'); setfont(rb, bold=True, color=RED)
rc = p.add_run('（价格与CVD同步创新低=空头趋势确认、非背离；资费自0.0093回落但仍为正=多头未投降、下行空间未尽；OI横盘58.5K=多头止损被新空替代。隔夜瀑布resume，6/3的超卖反抽thesis已失效、再创新低）'); setfont(rc)

line('盘面：', bold=True)
p = doc.add_paragraph()
setfont(p.add_run('日线级别　巨阴-3.96%破上升通道下轨+0.618斐波，跌穿日线POC 67,000、SMA200 74,885远在上方、全空排列　'))
setfont(p.add_run('利空'), bold=True, color=RED)
p = doc.add_paragraph()
setfont(p.add_run('小时级别　1H/4H全空排列、SMA20（66,335/69,406）向下并转头顶压力，15m跌破用户手绘支撑(6.3) 65,300-65,800　'))
setfont(p.add_run('利空'), bold=True, color=RED)

p = doc.add_paragraph()
setfont(p.add_run('趋势确认：'), bold=True)
setfont(p.add_run('空头趋势'), bold=True, color=RED)
setfont(p.add_run('（单边·瀑布破位下行）'))

p = doc.add_paragraph()
setfont(p.add_run('日内支撑：63,500-64,000　　日内压力：65,300-65,800'), bold=True)

line('策略：', bold=True)
line('中线主账户：6/3新开中线低多@67,000（50%仓·5x·700刀·左侧Taker·1:2.5·长持）隔夜随破位已止损出局·现无持仓（Notion状态=已平仓、按计划执行、未情绪化）。状态机＝现金观望、0%仓；空头单边未止跌前不新增中线多单、禁左侧接飞刀，待资费转负+CVD背离+超卖capitulation信号共振企稳再评估重建。')

doc.add_paragraph()
# ===== Maker =====
H('② Maker 计划（≥70 必挂 · 反弹空）')
line('BTC（策略）：执行概率 62%（已执行 0%）　胜率 62%', bold=True)
p = doc.add_paragraph()
setfont(p.add_run('反弹空：65,300-65,800'), bold=True, color=RED)
setfont(p.add_run('（跌破的用户支撑(6.3)转阻力 + 1H SMA20 66,335 下沿 = 最近且≥60%可达的高质量反抽阻力带，现价64,114需+1.8~2.6%）　持仓 700 刀　Binance　左侧　盈亏比 1:2.5 以上'))
line('止盈：63,400（TP1·+2.5R）减仓50%　/　62,700（TP2·+3.4R）')
line('止损：站稳 66,400 上方（有效突破1H SMA20 66,335+前结构=反抽空失效；单笔约1.3%，<3%）')
line('浮盈防守计划：', bold=True)
line('　第一层：价格首次触及 64,300（+1.5R）　设立4小时级别低点为止损、取消所有加仓')
line('　第二层：价格触及 63,400（+2.5R）　减仓50% 同时设保本损')
p = doc.add_paragraph()
setfont(p.add_run('注意事项：'), bold=True, color=ORANGE)
setfont(p.add_run('若价格优先触及 63,400（下行目标自行兑现、未反抽至入场区）则此策略取消（撤未成交挂单）'))
line('核心（开仓理由）：隔夜瀑布破位、价格与CVD同步创新低=空头趋势确认，多头未投降（资费仍正）下行空间足；价反抽至65,300-65,800（跌破的支撑6.3转阻力/1H SMA20下沿）遇压顺势做空。单挂边缘等待、价不到不成交（绝不追价）。')
p = doc.add_paragraph()
setfont(p.add_run('⚠️ 执行概率＝价格击中挂单区间的概率（限价单成交概率）·宪法≥60%。'), bold=True, color=ORANGE, size=9)
setfont(p.add_run('现价64,114距入场下沿65,300约+1.8%、距上沿65,800约+2.6%，former support flip+SMA20共振有反抽燃料→可达≈62%。取消价(63,400·成交前撤单)≠止损(66,400·成交后离场)。'), color=RGBColor(0x55,0x55,0x55), size=9)
line('持仓事实校准：主账户6/3中线低多@67,000隔夜已止损出局（无持仓）；Maker 6/3做多800刀已TP1止盈出局（盈$10），当前 Maker 无实际持仓。', color=RGBColor(0x55,0x55,0x55), size=9)

doc.add_paragraph()
# ===== 后附 =====
H('③ 后附（系统内部备份·不进 Notion 正文）')
line('评分明细（总分 71 = 盘口50 + 盘面21）', bold=True)
line('盘口 50/70：四维=CVD↓(创新低-34.4K)/价↓(创新低64,070)/资费下降但为正(0.0049自0.0093回落)/OI横盘(58.5K略升)=价与CVD同步创新低的「空头趋势确认」组合(无背离·强空)；多头未capitulate(资费仍正)属趋势中段。未给更高分因隔夜已-3.96%、24h多单爆仓8.04亿逼近超卖、反弹反包风险上升→50。')
line('盘面 21/30：接近度5（价刚破支撑6.3、距最近反弹阻力65,300-65,800约+1.9~2.6%、破位中非贴边）+流动性掠夺6（24h多单爆仓8.04亿、多头流动性大量被拿取）+SMA20 5（1H/4H/日全在下方且向下,满分）+量价3（放量下跌Vol456M配合,满分）+斐波趋势线2（破上升通道下轨+0.618,满分）。')
line('总分 71 ≥70 → 必挂 Maker（方向=空，挂价=高质量空头边缘且≥60%可达=65,300-65,800反抽带）。', bold=True)
line('五源支撑', bold=True)
line('① BlockBeats：利空主导=周四开门黑BTC破6.5万(现报64,609·24h-2.8%)/24h全网爆仓9.84亿(多单8.04亿·capitulation延续)/美股三大指数收跌(道指-1.2%/标普-0.7%/纳指-0.89%·微软英伟达跌逾3%)risk-off传导/日韩股市低开/美股资金分化270亿流入科技40亿流出其余=币市失血/Hyperliquid最大BTC空头21连胜获利332万/贝莱德存6005BTC潜在卖压。对冲(缓和)=特朗普对伊谈判"非常顺利"协议本周末或达成+伊朗四阶段首阶段全面停火(但今晨仍打击美军舰·未落实)/某巨鲸买401BTC逢低/Bitmine仿Strategy发3亿优先股囤币。噪音=WLD+39%等山寨AI逆势涨抽血BTC。中性偏空、部分兑现。')
line('② 日线footprint(Exo)：6/3完整Vol456M/Delta-27M(负)/CVD-112M(深负·自-85M续创低)/POC66,774(自69,354下移)/H-L3,424.3；现价64,114破POC下方。卖方完全主导、POC逐级下移、CVD累积深负=强空延续。')
line('③ Coinglass 1h(Bybit)：价64,070↓创新低/CVD-34,429K↓创新低/资金费率0.0049(正·自0.0093回落但仍正=多头未投降)/OI58.519K(横盘略升=新空进场+多头未平)=价与CVD同步创新低=趋势确认(非背离)·下跌中继。')
line('④ FedWatch：6/17不变98.4%、降息1.6%、加息0%；降息概率1月前6.7%→现1.6%(月内大幅price out)，1日前2.4%→1.6%微降→macro中性偏紧。')
line('⑤ TradingView 日/4H/1H/15m：日线收64,114(-3.96%·开66,760高67,516低64,092)破通道下轨+0.618+POC67,000，SMA200 74,885远上；4H64,109全空·SMA20 69,406向下·POC67,000头顶；1H64,106自68k砸穿67k POC一路下行·SMA20 66,335转阻力；15m64,130跌破用户支撑(6.3)65,300-65,800、压力(6.3)67,900-68,000与压力(6.2)71,500全在上方。最近高质量反抽空边缘=65,300-65,800。两层收口利空。')
line('不能做推导', bold=True)
line('空头单边+价砸至新低64,114→最易被收割侧=64k新低区右侧追空(隔夜瀑布、24h多单已爆8.04亿逼近capitulation，再创新低后追空易遭短挤/反弹反包)。禁「追做空」(仅禁右侧)；反抽至65,300-65,800的左侧反应空(=本Maker)、企稳后的左侧低多反包均不禁。')

out = r'C:\Users\31702\Desktop\Trading-System-Code\outputs\2026-06-04_早间日报.docx'
doc.save(out)
print('SAVED', out)
