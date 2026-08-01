# -*- coding: utf-8 -*-
"""Render a combined Markdown daily report into a clean DOCX.

The combined Markdown file is the single content source.  This renderer:
- converts Markdown headings to real Word heading styles;
- converts ``-`` items to real Word list paragraphs;
- converts ``**bold**`` and inline-code markers to formatted runs;
- reuses ``report_colors`` as the sole sentiment-color authority; and
- never exposes generator/debug labels in the document title.

Usage: py render_from_combined.py <combined_md_path> --date YYYY-MM-DD --slot SLOT
"""
import argparse
import hashlib
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import report_colors as rc  # 单一配色权威
from report_structure_gate import assert_combined_structure, assert_rendered_docx_structure
from runtime.reasoning_readiness import validate_locked_report_for_output
from runtime.state_store import atomic_json_write
from runtime.usage_logger import append_usage

# Notion 色名 → docx RGB（与现有手打脚本同值，统一来源）
DOCX_RGB = {
    "red": RGBColor(0xC0, 0x00, 0x00), "green": RGBColor(0x1E, 0x7A, 0x32),
    "yellow": RGBColor(0xB0, 0x60, 0x00), "purple": RGBColor(0x6A, 0x00, 0x99),
}
HEADING_BLUE = RGBColor(0x1F, 0x49, 0x7D)
SUBTITLE_GRAY = RGBColor(0x66, 0x66, 0x66)
COLOR_WORDS = {
    *rc.SENTIMENT.keys(),
    *rc.PREFIX_RED,
    *rc.PREFIX_GREEN,
    *rc.PREFIX_PURPLE,
}
SENT_WORDS = sorted(COLOR_WORDS, key=len, reverse=True)  # 长词优先(中性偏空 在 中性 之前)
LABELS = ("不能做", "基本面", "技术面", "盘口（合约）", "解读", "盘面", "日线级别",
          "小时级别", "趋势确认", "日内支撑", "日内压力", "策略", "BTC（策略）",
          "观望理由", "触发门槛", "核心", "情形一", "情形二", "证据完整性",
          "推理完整性", "交付完整性", "证据事实", "判断价值", "使用边界", "基准锚",
          "真实落值", "机械校验", "方向主判", "最强反证", "失效路径", "成交情景",
          "仓位账本")
INLINE_MARKUP = re.compile(r"(\*\*.*?\*\*|`[^`]*`)")
SLOT_CN = {"morning": "早间", "midday": "午间", "evening": "晚间"}
TITLE_FORBIDDEN = re.compile(
    r"(?i)(?:\.md\b|_combined\b|_render\b|\bmorning\b|\bmidday\b|"
    r"\bevening\b|\breport\b|\bv\d+\b|\bpoc\b|\bdebug\b|\btest\b)"
)
SOURCE_SUBHEADING = re.compile(
    r"^\d+\.\s+(?:BlockBeats|Coinglass|Exocharts|FedWatch|Farside|TradingView)\s*｜"
)
CRITICAL_PRICE_BLOCK = re.compile(
    r"^(?:策略：.*入场区间|止盈：|止损：|第一层：|第二层：|注意事项：|取消价≠止损：)"
)
REJECTED_CANDIDATE_AUDIT = re.compile(
    r"^被拒候选(?:否决原因|审计输入|风险情景|机械结论)（不可执行，不得用于下单）"
)
CRITICAL_PRICE_CHAINS = (
    (
        re.compile(r"^策略：.*入场区间"),
        re.compile(r"^止盈："),
        re.compile(r"^止损："),
    ),
    (
        re.compile(r"^第一层："),
        re.compile(r"^第二层："),
        re.compile(r"^注意事项："),
        re.compile(r"^取消价≠止损："),
    ),
)


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_docx_path(trading_date, slot, root=ROOT):
    return Path(root) / "outputs" / f"{trading_date}_{slot}_render.docx"


def render_manifest_path(trading_date, slot, root=ROOT):
    return Path(root) / "runtime_state" / trading_date / f"render_{slot}.json"


def _load_object(path):
    if not Path(path).is_file():
        return None
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"{path} must contain a JSON object")
    return payload


def preflight(lines, trading_date, slot):
    """Resolve a safe title and block known render-loop markers before DOCX."""
    if slot not in SLOT_CN:
        raise ValueError(f"unsupported slot: {slot}")
    lone_markers = [
        index
        for index, line in enumerate(lines, 1)
        if line.strip() == ">"
    ]
    if lone_markers:
        raise ValueError(f"lone blockquote marker at line(s): {lone_markers}")
    title = assert_combined_structure(lines, trading_date, slot)
    forbidden = TITLE_FORBIDDEN.search(title)
    if forbidden:
        raise ValueError(f"forbidden generator/debug marker in title: {forbidden.group(0)}")
    return title


def renderer_sha256():
    digest = hashlib.sha256()
    for path in (
        Path(__file__),
        Path(__file__).with_name("report_colors.py"),
        Path(__file__).with_name("report_structure_gate.py"),
        ROOT / "governance" / "vnext_report_presentation_contract.json",
    ):
        digest.update(path.name.encode("utf-8"))
        digest.update(path.read_bytes())
    return digest.hexdigest()


def setfont(run, size=10.5, bold=False, color=None):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def color_for(word):
    c = rc.expected_color(word)
    return DOCX_RGB.get(c) if c else None


def plain_markdown(text):
    """Remove supported inline Markdown markers without changing content."""
    return re.sub(r"\*\*(.*?)\*\*", r"\1", text).replace("`", "")


def _emit_colored_text(
    paragraph,
    text,
    size,
    force_bold=False,
    code=False,
    neutral_audit=False,
):
    for tok in re.split(
        "(" + "|".join(map(re.escape, SENT_WORDS)) + ")",
        text,
    ):
        if not tok:
            continue
        sentiment_colored = not neutral_audit and tok in COLOR_WORDS
        col = SUBTITLE_GRAY if neutral_audit else color_for(tok) if sentiment_colored else None
        run = paragraph.add_run(tok)
        setfont(
            run,
            size=size,
            bold=force_bold or sentiment_colored,
            color=col,
        )
        if code:
            run.font.name = "Consolas"


def emit_paragraph(doc, text, size=10.5, style=None):
    """Write one Markdown line as formatted Word runs."""
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    semantic_text = plain_markdown(text).strip()
    neutral_audit = bool(REJECTED_CANDIDATE_AUDIT.match(semantic_text))
    idx = 0
    m = re.match(r"^(" + "|".join(map(re.escape, LABELS)) + r")(：|:)?", text)
    if m:
        r = p.add_run(m.group(0))
        setfont(r, size=size, bold=True)
        idx = m.end()
    rest = text[idx:]
    cursor = 0
    for markup in INLINE_MARKUP.finditer(rest):
        if markup.start() > cursor:
            _emit_colored_text(
                p,
                rest[cursor:markup.start()],
                size,
                neutral_audit=neutral_audit,
            )
        token = markup.group(0)
        if token.startswith("**"):
            _emit_colored_text(
                p,
                token[2:-2],
                size,
                force_bold=True,
                neutral_audit=neutral_audit,
            )
        else:
            _emit_colored_text(
                p,
                token[1:-1],
                size,
                code=True,
                neutral_audit=neutral_audit,
            )
        cursor = markup.end()
    if cursor < len(rest):
        _emit_colored_text(
            p,
            rest[cursor:],
            size,
            neutral_audit=neutral_audit,
        )
    if SOURCE_SUBHEADING.match(semantic_text):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
    if CRITICAL_PRICE_BLOCK.match(semantic_text):
        p.paragraph_format.keep_together = True
    if neutral_audit:
        p.paragraph_format.keep_together = True
        # Keep only the reason heading paired with the next audit paragraph.
        # Risk scenarios can reach 31 rows for a five-level ladder and must
        # remain independently pageable instead of forming one unbounded chain.
        if semantic_text.startswith("被拒候选否决原因"):
            p.paragraph_format.keep_with_next = True
    return p


def _apply_critical_price_keep_chains(doc):
    """Keep only complete adjacent critical-price sequences on one page flow."""

    paragraphs = list(doc.paragraphs)
    for chain in CRITICAL_PRICE_CHAINS:
        width = len(chain)
        for start in range(0, len(paragraphs) - width + 1):
            group = paragraphs[start : start + width]
            if not all(
                pattern.match(paragraph.text.strip())
                for pattern, paragraph in zip(chain, group, strict=True)
            ):
                continue
            for index, paragraph in enumerate(group):
                paragraph.paragraph_format.keep_together = True
                if index < width - 1:
                    paragraph.paragraph_format.keep_with_next = True


def render(combined_path, out_path, *, source_title=None):
    with open(combined_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Microsoft YaHei'
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.0
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.right_margin = Inches(1.25)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
    for style_name, size in (("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = doc.styles[style_name]
        style.font.name = 'Microsoft YaHei'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = HEADING_BLUE
        style.paragraph_format.space_before = Pt(6 if style_name != "Heading 3" else 2)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.keep_with_next = True
    list_style = doc.styles["List Bullet"]
    list_style.font.name = 'Microsoft YaHei'
    list_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    list_style.font.size = Pt(10.5)
    list_style.paragraph_format.space_after = Pt(2)
    list_style.paragraph_format.line_spacing = 1.0

    base = os.path.basename(combined_path).replace("_combined.md", "")
    display_base = (
        base.replace("_evening", " 晚间")
        .replace("_morning", " 早间")
        .replace("_midday", " 午间")
    )
    source_title = source_title or next(
        (
            plain_markdown(line[2:].strip())
            for line in lines
            if line.startswith("# ")
        ),
        f"{display_base}日报",
    )
    title_main = source_title
    title_detail = ""
    if "（" in source_title and source_title.endswith("）"):
        title_main, detail = source_title.split("（", 1)
        title_detail = "（" + detail

    hp = doc.add_paragraph()
    hp.paragraph_format.space_after = Pt(3 if title_detail else 6)
    hp.paragraph_format.keep_with_next = True
    hr = hp.add_run(title_main.strip())
    setfont(hr, size=15, bold=True, color=HEADING_BLUE)
    if title_detail:
        subtitle = doc.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(10)
        sr = subtitle.add_run(title_detail)
        setfont(sr, size=10.5, color=SUBTITLE_GRAY)

    colored = 0
    for ln in lines:
        s = ln.rstrip()
        if not s:
            continue
        if s == ">":                                  # 合法 quote 段内空行
            continue
        if s.startswith("# "):
            continue
        if s.startswith("> "):                       # 作者自查 → 小号注释
            emit_paragraph(doc, s[2:], size=9)
        elif s.startswith("#### "):
            raise ValueError("H4 is forbidden by the gold-template fidelity contract")
        elif s.startswith("### "):
            label = plain_markdown(s[4:])
            if re.match(r"^[A-D]\.\s", label):
                p = doc.add_paragraph(style="Heading 2")
                r = p.add_run(label)
                setfont(r, size=12, bold=True, color=HEADING_BLUE)
            else:
                p = doc.add_paragraph(style="Heading 3")
                p.paragraph_format.space_before = Pt(2)
                r = p.add_run(label)
                setfont(r, size=10.5, bold=True)
        elif s.startswith("## "):                     # 段标题(①②③)
            p = doc.add_paragraph(style="Heading 1")
            label = plain_markdown(s[3:])
            r = p.add_run(label)
            if label == "不能做":
                setfont(r, size=10.5, bold=True)
                p.paragraph_format.space_before = Pt(2)
            else:
                setfont(r, size=13, bold=True, color=HEADING_BLUE)
        elif re.match(r"^\s*-\s+", s):                # ③后附/情形条目
            item = re.sub(r"^\s*-\s+", "", s)
            emit_paragraph(doc, item, style="List Bullet")
        else:
            emit_paragraph(doc, s)

    _apply_critical_price_keep_chains(doc)

    # 统计上色 run 数(验证用)
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.color and r.font.color.rgb in DOCX_RGB.values():
                colored += 1

    doc.save(out_path)
    return len(doc.paragraphs), colored


def render_canonical(
    combined_path,
    trading_date,
    slot,
    *,
    root=ROOT,
    vnext_run_root=None,
    _vnext_lock_acquired=False,
):
    root = Path(root).resolve()
    source = Path(combined_path).resolve()
    if vnext_run_root is not None and not _vnext_lock_acquired:
        from runtime_vnext.artifacts import ArtifactStore

        guarded_run_root = Path(vnext_run_root).resolve()
        store = ArtifactStore(root)
        with store.controller_lock(guarded_run_root):
            guarded_run = store.load_run(guarded_run_root)
            if (
                guarded_run.get("schema_version") != "daily-run-vnext-v2"
                or guarded_run.get("run_mode") not in {"shadow", "production"}
            ):
                raise ValueError(
                    "legacy vNext runs are read-only; rendering requires an explicitly authorized v2 run mode"
                )
            store._assert_active(guarded_run_root)
            return render_canonical(
                source,
                trading_date,
                slot,
                root=root,
                vnext_run_root=guarded_run_root,
                _vnext_lock_acquired=True,
            )
    run_root = None
    if vnext_run_root is None:
        validate_locked_report_for_output(trading_date, slot, source, root=root)
    else:
        from runtime_vnext.artifacts import ArtifactStore

        run_root = Path(vnext_run_root).resolve()
        expected_run_root = root / "runtime_state_vnext" / trading_date / slot
        run_root.relative_to(expected_run_root.resolve())
        vnext_store = ArtifactStore(root)
        guarded_run = vnext_store.load_run(run_root)
        if (
            guarded_run.get("schema_version") != "daily-run-vnext-v2"
            or guarded_run.get("run_mode") not in {"shadow", "production"}
        ):
            raise ValueError(
                "legacy vNext runs are read-only; rendering requires an explicitly authorized v2 run mode"
            )
        certified = vnext_store.load_stage_output_record(
            run_root,
            "report",
            "combined",
        )
        certified_source = (root / certified["path"]).resolve()
        if source != certified_source:
            raise ValueError("vNext renderer source differs from certified report output")
    lines = source.read_text(encoding="utf-8").splitlines()
    title = preflight(lines, trading_date, slot)
    if run_root is None:
        target = canonical_docx_path(trading_date, slot, root).resolve()
        manifest_path = render_manifest_path(trading_date, slot, root)
        log_usage = append_usage
    else:
        delivery_assets = run_root / "delivery_assets"
        target = (
            delivery_assets / f"{trading_date}_{slot}_render.docx"
        ).resolve()
        manifest_path = delivery_assets / "render_manifest.json"
        log_usage = lambda **_kwargs: None
    duplicates = [
        path
        for path in target.parent.glob(f"{trading_date}*.docx")
        if path.resolve() != target
        and (
            f"_{slot}" in path.stem.lower()
            or SLOT_CN[slot] in path.stem
        )
    ]
    if duplicates:
        raise ValueError(
            "non-canonical DOCX artifacts exist for this report: "
            + ", ".join(path.name for path in duplicates)
        )

    source_hash = sha256_file(source)
    renderer_hash = renderer_sha256()
    presentation_contract_hash = sha256_file(
        root / "governance" / "vnext_report_presentation_contract.json"
    )
    existing = _load_object(manifest_path) or {}
    key = hashlib.sha256(
        f"{trading_date}:{slot}:{source_hash}:{renderer_hash}".encode("ascii")
    ).hexdigest()
    if (
        existing.get("status") == "success"
        and existing.get("source_sha256") == source_hash
        and existing.get("renderer_sha256") == renderer_hash
        and target.is_file()
        and existing.get("docx_sha256") == sha256_file(target)
    ):
        log_usage(
            trading_date=trading_date,
            slot=slot,
            phase="render_docx",
            status="success",
            event_id=f"{key}:render:cache",
            retries=max(int(existing.get("attempts", 1)) - 1, 0),
            details={
                "cache_hit": True,
                "attempt": existing.get("attempts", 1),
                "backend": "python-docx",
                "input_sha256": source_hash,
                "renderer_sha256": renderer_hash,
                "output_sha256": existing["docx_sha256"],
                "docx_path": str(target),
            },
            root=root,
        )
        result = dict(existing)
        result["cache_hit"] = True
        return result

    attempts = int(existing.get("attempts", 0))
    if attempts >= 2:
        log_usage(
            trading_date=trading_date,
            slot=slot,
            phase="render_docx",
            status="blocked",
            event_id=f"{key}:render:attempt-limit",
            retries=1,
            details={
                "cache_hit": False,
                "attempt": attempts,
                "backend": "python-docx",
                "input_sha256": source_hash,
                "renderer_sha256": renderer_hash,
                "failure_reason": "render attempt limit reached; third attempt forbidden",
            },
            root=root,
        )
        raise RuntimeError(
            "render attempt limit reached (2); third attempt is forbidden"
        )
    attempt = attempts + 1
    started = time.perf_counter()
    temporary = target.with_name(f".{target.name}.attempt-{attempt}.tmp.docx")
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        paragraph_count, colored_runs = render(
            source,
            temporary,
            source_title=title,
        )
        assert_rendered_docx_structure(temporary, trading_date, slot)
        os.replace(temporary, target)
        docx_hash = sha256_file(target)
        payload = {
            "schema_version": "1.0",
            "status": "success",
            "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "trading_date": trading_date,
            "slot": slot,
            "attempts": attempt,
            "cache_hit": False,
            "backend": "python-docx",
            "source_path": str(source),
            "source_sha256": source_hash,
            "renderer_sha256": renderer_hash,
            "presentation_contract_sha256": presentation_contract_hash,
            "resolved_title": title,
            "docx_path": str(target),
            "docx_sha256": docx_hash,
            "paragraph_count": paragraph_count,
            "colored_runs": colored_runs,
        }
        atomic_json_write(manifest_path, payload)
        log_usage(
            trading_date=trading_date,
            slot=slot,
            phase="render_docx",
            status="success",
            event_id=f"{key}:render:{attempt}",
            duration_ms=round((time.perf_counter() - started) * 1000),
            retries=attempt - 1,
            details={
                "cache_hit": False,
                "attempt": attempt,
                "backend": "python-docx",
                "input_sha256": source_hash,
                "renderer_sha256": renderer_hash,
                "output_sha256": docx_hash,
                "docx_path": str(target),
            },
            root=root,
        )
        return payload
    except BaseException as exc:
        temporary.unlink(missing_ok=True)
        failure = {
            "schema_version": "1.0",
            "status": "failed",
            "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "trading_date": trading_date,
            "slot": slot,
            "attempts": attempt,
            "source_path": str(source),
            "source_sha256": source_hash,
            "renderer_sha256": renderer_hash,
            "presentation_contract_sha256": presentation_contract_hash,
            "docx_path": str(target),
            "failure_reason": str(exc),
        }
        atomic_json_write(manifest_path, failure)
        log_usage(
            trading_date=trading_date,
            slot=slot,
            phase="render_docx",
            status="failed",
            event_id=f"{key}:render:{attempt}:failed",
            duration_ms=round((time.perf_counter() - started) * 1000),
            retries=attempt - 1,
            details={
                "cache_hit": False,
                "attempt": attempt,
                "backend": "python-docx",
                "input_sha256": source_hash,
                "renderer_sha256": renderer_hash,
                "failure_reason": str(exc),
            },
            root=root,
        )
        raise


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("combined", type=Path)
    parser.add_argument("--date", required=True)
    parser.add_argument("--slot", required=True, choices=tuple(SLOT_CN))
    parser.add_argument("--vnext-run-id")
    args = parser.parse_args()
    vnext_run_root = (
        ROOT / "runtime_state_vnext" / args.date / args.slot / args.vnext_run_id
        if args.vnext_run_id
        else None
    )
    result = render_canonical(
        args.combined,
        args.date,
        args.slot,
        vnext_run_root=vnext_run_root,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))
