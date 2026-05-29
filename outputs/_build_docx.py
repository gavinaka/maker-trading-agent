# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# 中文字体设置
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

# 页边距
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color

def add_h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 18, True, RGBColor(0x1F, 0x3A, 0x5F))

def add_h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, 13, True, RGBColor(0x1F, 0x3A, 0x5F))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)

def add_h3(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, 11.5, True, RGBColor(0x2C, 0x5A, 0x8C))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)

def add_para(segments):
    """segments = list of (text, bold) tuples"""
    p = doc.add_paragraph()
    for text, bold in segments:
        r = p.add_run(text)
        set_font(r, 10.5, bold)
    return p

def add_quote(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run(text)
    set_font(r, 9.5, False, RGBColor(0x66, 0x66, 0x66))
    r.italic = True

def add_bullet(text_segments):
    p = doc.add_paragraph(style='List Bullet')
    for text, bold in text_segments:
        r = p.add_run(text)
        set_font(r, 10.5, bold)

def add_hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pBdr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single', qn('w:sz'): '6',
        qn('w:space'): '1', qn('w:color'): 'BFBFBF'
    })
    pBdr.append(bottom)
    pPr.append(pBdr)

# ===== 标题 =====
add_h1('主账户日报')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('2026-05-29  |  Trading-System v2026.5  |  output_agent 渲染')
set_font(r, 10, False, RGBColor(0x88, 0x88, 0x88))

add_quote('渲染来源：templates/main_account_daily.md（图谱第十节）')
add_quote('数据来源：data_sources/inbox/2026-05-29/、state/main_account_state.json')
add_quote('说明：盘口/盘面字段因当日 Exocharts、Coinglass、FedWatch 暂未采集，按规则标注 [待补充]，不删除字段、不省略段落。')

add_hr()

# ===== 日报正文 =====
add_h2('一、日报正文')

# 日期
p = doc.add_paragraph()
r = p.add_run('2026-05-29')
set_font(r, 12, True)

# 不能做
add_para([('不能做：', True), ('做多', False)])

# 基本面
add_para([
    ('基本面：', True),
    ('中性（黑客事件与机构存入交易所 vs Coinbase 合规扩张相抵；美联储通胀信号矛盾：截尾平均 2.3% 偏冷、核心 PCE 3.3% 偏热）', False),
    ('价格未兑现', False),
])

# 重要新闻
add_h3('重要新闻（含时间戳，来源 BlockBeats 截图 ss_4739pehi4）')
news_items = [
    ('[10:43]', '纽交所母公司 ICE CEO：Hyperliquid 比纳斯达克还大，正在交流', 'blockbeats_001'),
    ('[10:39]', '派盾：DxSale 遭攻击，BSC 上旧 LP 被抽取 730 万美元', 'blockbeats_001'),
    ('[10:29]', 'Ventuals 回应 SPACEX-USDH 暴跌：预言机错误数据致异常清算，48h 内赔付', 'blockbeats_001'),
    ('[10:21]', '沃什偏爱通胀指标跌至 2.3%，但核心 PCE 升至 3.3%（2023 年以来最快）', 'blockbeats_001'),
    ('[10:11]', 'Strategy 向 Coinbase Prime 存入 411.48 枚 BTC（3030 万美元）', 'blockbeats_001'),
    ('[09:47]', 'Coinbase 推出 TRUSThub 全球加密合规网络，BNY Mellon 加入', 'blockbeats_001'),
]
for ts, content, shot in news_items:
    add_bullet([(ts + ' ', True), (content, False), ('  → 截图：' + shot, False)])

add_hr()

# 技术面
add_h2('二、技术面')

add_h3('盘口（合约）')
add_para([
    ('CVD ', True), ('[待补充 - Exocharts/Coinglass 当日未采集]　', False),
    ('价格 ', True), ('[待补充 - Exocharts/Coinglass 当日未采集]　', False),
    ('资金费率 ', True), ('[待补充 - Coinglass 当日未采集]　', False),
    ('持仓 ', True), ('[待补充 - Coinglass 当日未采集]', False),
])
add_para([('解读：', True), ('[待补充 - 盘口 4 项缺失，无法进入 footprint 16 组合判定]', False)])

add_h3('盘面')
add_para([('日线级别 ', True), ('[待补充 - exocharts_btc_1d.json 未生成]　[待补充]', False)])
add_para([('小时级别 ', True), ('[待补充 - exocharts_btc_1h.json 未生成]　[待补充]', False)])

add_hr()

# 趋势 & 支撑压力
add_h2('三、趋势确认与日内支撑压力')
add_para([('趋势确认：', True), ('空头趋势（中短线）', False)])
add_para([
    ('日内支撑：', True), ('[待补充 - 1h 关键位未采集]　　', False),
    ('日内压力：', True), ('[待补充 - 1h 关键位未采集]', False),
])

add_hr()

# 策略
add_h2('四、策略')
add_para([('持中线做空', True), ('（5 层防守，已激活保本损 + 动态止损）', False)])
add_bullet([('持仓：', True), ('BTCUSDT short / 均价 79,150 / 仓位 $500 / 5x / 50%', False)])
add_bullet([('中线计划：', True), ('保持不变（active_plan = true，layer 5/5，breakeven_active = true）', False)])
add_bullet([('日内：', True), ('维持原计划，不新增仓位；Maker 仅允许做空方向（受主账户 cannot_do = "做多" 约束）', False)])
add_bullet([('风控：', True), ('单笔止损 ≤ 3%、5 项状态门禁全绿（情绪平静 / 节奏一致 / 无结果依赖 / 连续止损 0 / 身体正常）', False)])

add_hr()

# ===== 数据完整度 =====
add_h2('五、数据完整度与下一步')

table = doc.add_table(rows=1, cols=3)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
for i, t in enumerate(['数据源', '状态', '影响段落']):
    hdr[i].text = ''
    p = hdr[i].paragraphs[0]
    r = p.add_run(t)
    set_font(r, 10.5, True, RGBColor(0xFF, 0xFF, 0xFF))

rows_data = [
    ('BlockBeats', '✅ 已采集（ss_4739pehi4）', '基本面 / 重要新闻'),
    ('Exocharts (BTC 1d/1h)', '❌ 未采集', '盘口 CVD/价格、盘面日/小时、支撑压力'),
    ('Coinglass (合约)', '❌ 未采集', '盘口 资金费率 / 持仓'),
    ('CME FedWatch', '❌ 未采集', '基本面（仅文字提及，未量化）'),
]
for row in rows_data:
    cells = table.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = ''
        p = cells[i].paragraphs[0]
        r = p.add_run(val)
        set_font(r, 10)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
r = p.add_run('下一步建议：')
set_font(r, 10.5, True)
r = p.add_run('补采 Exocharts + Coinglass + FedWatch → 重跑 footprint_agent + scoring_agent → 在 11:00 评分点产出完整 70+30 评分，再决定 Maker 上午是否触发开仓。')
set_font(r, 10.5)

add_hr()

# Footer
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Generated by output_agent  |  Template: main_account_daily.md v2026.5')
set_font(r, 9, False, RGBColor(0x99, 0x99, 0x99))
r.italic = True

out_path = r'C:\Users\31702\Desktop\Trading-System\outputs\2026-05-29_主账户日报.docx'
doc.save(out_path)
print('SAVED:', out_path)
