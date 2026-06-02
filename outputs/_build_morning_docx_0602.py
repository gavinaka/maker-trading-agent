# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
# base font
st = doc.styles['Normal']
st.font.name = 'Microsoft YaHei'
st.font.size = Pt(10.5)
st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def setfont(run, size=10.5, bold=False, color=None):
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

def H(text, size=13):
    p = doc.add_paragraph()
    r = p.add_run(text)
    setfont(r, size=size, bold=True, color=RGBColor(0x1F,0x49,0x7D))
    p.space_after = Pt(4)
    return p

def line(text, bold=False, color=None, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    setfont(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(2)
    return p

RED = RGBColor(0xC0,0x00,0x00)
ORANGE = RGBColor(0xB0,0x60,0x00)
GREEN = RGBColor(0x1E,0x7A,0x32)

# Title
p = doc.add_paragraph()
r = p.add_run('2026-06-02 早间日报（系统独立稿）')
setfont(r, size=16, bold=True, color=RGBColor(0x00,0x00,0x00))
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

line('门禁 6/2＝可交易（身体正常/情绪平静/节奏顺/连续止损0）；心力6/1无异常·专注5→正常，无收缩机制。总评分 72（盘口50+盘面22）≥70 → 必挂 Maker 单于高质量边缘。', color=RGBColor(0x55,0x55,0x55), size=9)

# ===== 正文 =====
H('① 日报正文')
line('不能做：（追）做空', bold=True, color=RED)
p = doc.add_paragraph()
r1 = p.add_run('基本面：'); setfont(r1, bold=True)
r2 = p.add_run('中性偏空'); setfont(r2, bold=True, color=ORANGE)
r3 = p.add_run('（地缘缓和 Tier1 利多——特朗普推美伊和谈"一周内或达成"+黎以停火、油价涨7%后回落——未提振币市；加密结构性去杠杆主导：Strategy 首披露抛售32枚BTC打破"永不卖"、加密概念股普跌CRCL/BLSH跌超7%·MSTR-5.85%、24h爆仓6.4亿主爆多单、恐慌指数23极度恐慌、巨鲸抛/做空ETH）'); setfont(r3)
r4 = p.add_run('价格已部分兑现'); setfont(r4, bold=True)
r5 = p.add_run('（BTC已跌至71,285）'); setfont(r5)

line('技术面：', bold=True)
line('盘口（合约）：', bold=True)
line('CVD下降　价格下降　资金费率上升（为正值）　持仓下降')
p = doc.add_paragraph()
ra = p.add_run('解读：'); setfont(ra, bold=True)
rb = p.add_run('中性偏空'); setfont(rb, bold=True, color=ORANGE)
rc = p.add_run('（放量破位+主动卖压持续；OI下降=多头去杠杆离场而非空头新增、资金费率仍正=残余抄底盘未投降，下跌中继/多头出清，方向偏空但非空头强增仓）'); setfont(rc)

line('盘面：', bold=True)
p = doc.add_paragraph()
setfont(p.add_run('日线级别　自约82,000回落跌破SMA20(75,960)+0.618，POC逐级下移，6/1放量破位(Vol 261M/CVD-38M/H-L 3394.9)，上升通道下沿测试中　'))
setfont(p.add_run('利空'), bold=True, color=RED)
p = doc.add_paragraph()
setfont(p.add_run('小时级别　价于1H SMA20(71,888)下方、下降趋势线压制，贴71,000多头最后防守，CVD↓　'))
setfont(p.add_run('利空'), bold=True, color=RED)

p = doc.add_paragraph()
setfont(p.add_run('趋势确认：'), bold=True)
setfont(p.add_run('空头趋势'), bold=True, color=RED)
setfont(p.add_run('（短线/中短线·放量破位下行）'))

p = doc.add_paragraph()
setfont(p.add_run('日内支撑：70,600-71,000　　日内压力：71,800-72,300'), bold=True)

line('策略：', bold=True)
line('中线多单（5/28做多@73,500·底仓30%·5x·$1000·1:3）保持不变、死守 71,000（上升通道下轨/有效跌破即止损出局）；现价 71,285 浮亏 -3.0%、距止损 285 点高度戒备；浮盈防守层(75,500/76,200/74,400)均未激活、不加仓不止盈。框架（中线＝基本面>盘面>盘口）：基本面中性偏空+盘面空头+价临界，多单逻辑承压，但纪律守预设止损、未破不主观离场。')

doc.add_paragraph()
# ===== Maker =====
H('② Maker 计划（≥70 必挂）')
line('BTC（策略）：执行概率 ~60%（已执行 0%）　胜率 ~60%', bold=True)
p = doc.add_paragraph()
setfont(p.add_run('反弹空：71,800-72,000'), bold=True, color=RED)
setfont(p.add_run('（0.618 fib 71,824 + 1H SMA20 71,888 + 大级别共振支撑下沿 = 首压反抽阻力带，最近且≥60%可达的高质量空头边缘）　持仓 700 刀　Binance　左侧　盈亏比 1:2.5 以上'))
line('止盈：70,600（TP1·多头最后防守下沿）减仓50%　/　TP2 70,000-70,100')
line('止损：站稳 72,500 上方（反抽高点/0.618 上沿失守）= 结构失效离场（单笔约0.8%，<3%）')
line('浮盈防守计划：', bold=True)
line('　第一层：价格触及 71,000（1.5R·多头最后防守）以下　设立4小时级别极点（反抽高点）为止损 且取消所有加仓')
line('　第二层：价格触及 70,700（2R）以下　设立保本损')
p = doc.add_paragraph()
setfont(p.add_run('注意事项：'), bold=True, color=ORANGE)
setfont(p.add_run('若价格优先触及 70,600（先向下破多头最后防守、未反抽至入场区）则此策略取消（撤未成交挂单）'))
line('核心（开仓理由）：空头放量破位+0.618/1H SMA20/大级别共振支撑下沿构成首压反抽阻力带，超跌反弹至首压被拒做空、顺势空头延续')
p = doc.add_paragraph()
setfont(p.add_run('⚠️ 执行概率修正（用户拍板）：'), bold=True, color=ORANGE, size=9)
setfont(p.add_run('执行概率＝价格击中挂单区间的概率（限价单成交概率），非"我是否挂单"；宪法要求≥60%。原稿挂73,000-73,165距现价约+1,800点（逆势反抽可达≈35%·不达门槛），已下移至首压71,800-72,000反抽带（现价71,285仅需+0.7%~+1.0%、且处71,000支撑有反弹燃料，可达≈60%）。价未到→单挂边缘等待不成交（绝不追价）。取消价(70,600·成交前撤单)≠止损(72,500·成交后离场)。'), color=RGBColor(0x55,0x55,0x55), size=9)
line('持仓事实校准：Maker 另有 6/1 做多700刀 @71,000 持仓中（入场价由防守层1.5R=72,500、2R=73,000反推：R=1,000、入场71,000、隐含止损≈70,000；条目未显式写均价），现价71,285=+0.4%（约+0.3R）小幅浮盈；防守层72,500(1.5R)/73,000(2R)未激活、取消价73,000已失效。本反弹空挂71,800-72,000，与该多单不在同价位、并存不冲突。', color=RGBColor(0x55,0x55,0x55), size=9)

doc.add_paragraph()
# ===== 后附 =====
H('③ 后附（系统内部备份·不进 Notion 正文）')
line('评分明细', bold=True)
line('盘口 50/70：CVD↓+价↓+资金费率正(↑)+OI↓。核心放量卖压偏空成立；OI↓=多头去杠杆（非C01/C02空头强增仓）、费率仍正=多头未投降，增仓动能折扣 → 手动50。')
line('盘面 22/30：接近度8（价极贴71,000）+流动性掠夺4+SMA20 5（空头排列）+量价3（6/1放量261M）+斐波趋势线2（破0.618/下降趋势线下）。')
line('总分 72 ≥70 → 必挂 Maker（方向=空，挂价=高质量空头边缘且≥60%可达=71,800-72,000首压反抽带）。', bold=True)
line('五源支撑', bold=True)
line('① BlockBeats：Tier1美伊和谈推进+黎以停火（risk-on利多·油价涨7%回落）、美股创新高英伟达+6%；Tier2行业利空主导——Strategy抛32BTC、加密概念股-7%、爆仓6.4亿主爆多单、恐慌指数23、巨鲸抛/空ETH。利空5/利多2、地缘未提振→中性偏空，部分兑现。')
line('② 日线footprint：POC逐级下移→71082(6/2)，Delta持续负，CVD累积深负，6/1放量破位261M/CVD-38M/H-L3394=趋势性向下突破。')
line('③ Coinglass 1h：价71,263↓/CVD-13.36K↓/资金费率0.0097↑(正)/OI58.18K↓(自62k回落)=多头去杠杆离场+残余抄底盘未投降。')
line('④ FedWatch：6/17不变98.4%、降息1.6%微升、加息0%，较1日前几无变化→macro中性。')
line('⑤ TradingView：现价71,285（1H/15m=71,285·71,284、4H=71,290、日线=71,298，取最新最细粒度71,285）已破0.618(71,824)+共振(73,165)+SMA20，贴多头最后防守71,000。阻力梯71,824→73,165→73,932→74,250-75,300(空头流动性)→76,000；支撑70,600-71,000→70,400-70,634。两层收口利空。位置把关：追空在71,000上方非高质量→挂价上移至首压71,800-72,000反抽带（≥60%可达；更高73,000-73,165反抽距现价约+1,800点可达<60%不达门槛）。')
line('不能做推导', bold=True)
line('空头单边+价贴多头最后防守71,000+多头去杠杆(OI↓/爆仓6.4亿)→最易被收割侧=71,000支撑上方追空(右侧)。禁「追做空」(仅禁右侧)；对侧反弹左侧反应空(=本Maker)不禁、71,000掠夺后反包左侧做多不禁。本反弹空(左侧)≠追空(右侧)。')

out = r'C:\Users\31702\Desktop\Trading-System-Code\outputs\2026-06-02_早间日报.docx'
doc.save(out)
print('SAVED', out)
