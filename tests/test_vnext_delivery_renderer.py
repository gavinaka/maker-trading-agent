from __future__ import annotations

import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from docx import Document
from outputs import render_from_combined as canonical_report_renderer
from outputs.render_from_combined import render
from runtime.state_store import atomic_json_write
from runtime_vnext.artifacts import ArtifactStore
from runtime_vnext.delivery_renderer import (
    _artifact_record,
    _manifest_files_valid,
    WORD_QA_REVIEW_CHECKLIST_KEYS,
    WORD_QA_REVIEW_CHECKLIST_VERSION,
    approve_run_qa,
    delivery_control_state,
    prepare_run_qa,
    reject_run_qa,
)
from runtime_vnext.errors import FailureClass, WorkflowFailure
from runtime_vnext.methodology import canonical_sha256, sha256_file


class VNextDeliveryRendererTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.store = ArtifactStore(self.root)
        self.run_id = "run-word-qa-test"
        self.store.create_run(
            "2026-07-14",
            "midday",
            methodology_manifest={"manifest_sha256": "a" * 64},
            run_mode="shadow",
            authorization_ref="test-shadow-renderer-approval",
            run_id=self.run_id,
        )
        self.run_root = self.store.run_root("2026-07-14", "midday", self.run_id)

        report = self.run_root / "report" / "attempt-001" / "combined.md"
        report.parent.mkdir(parents=True)
        report.write_text("certified report", encoding="utf-8")
        report_record = _artifact_record(self.root, report)
        checkpoint = self.store.load_checkpoint(self.run_root)
        checkpoint["completed"] = {
            "report": {
                "attempt": 1,
                "manifest_path": "unused",
                "manifest_sha256": "b" * 64,
                "outputs": {"combined": report_record},
            }
        }
        checkpoint["next_stage"] = "delivery"
        checkpoint.pop("checkpoint_sha256", None)
        checkpoint["checkpoint_sha256"] = canonical_sha256(checkpoint)
        atomic_json_write(self.run_root / "checkpoint.json", checkpoint)

        presentation_contract = (
            self.root / "governance" / "vnext_report_presentation_contract.json"
        )
        presentation_contract.parent.mkdir(parents=True)
        presentation_contract.write_bytes(
            (
                Path(__file__).resolve().parents[1]
                / "governance"
                / "vnext_report_presentation_contract.json"
            ).read_bytes()
        )
        assets = self.run_root / "delivery_assets"
        assets.mkdir()
        self.docx = assets / "2026-07-14_midday_render.docx"
        self.docx.write_bytes(b"PK-test-docx")
        atomic_json_write(
            assets / "render_manifest.json",
            {
                "schema_version": "1.0",
                "status": "success",
                "trading_date": "2026-07-14",
                "slot": "midday",
                "attempts": 1,
                "source_path": str(report),
                "source_sha256": report_record["sha256"],
                "renderer_sha256": canonical_report_renderer.renderer_sha256(),
                "presentation_contract_sha256": sha256_file(
                    presentation_contract
                ),
                "docx_path": str(self.docx),
                "docx_sha256": sha256_file(self.docx),
            },
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _fake_export(self, docx: Path, pdf: Path, **kwargs) -> str:
        self.assertEqual(docx, self.docx.resolve())
        self.assertEqual(kwargs["trading_date"], "2026-07-14")
        self.assertEqual(kwargs["slot"], "midday")
        self.assertEqual(kwargs["root"], self.root)
        pdf.write_bytes(b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\nstartxref\n0\n%%EOF\n")
        return "word_com"

    def _remove_seeded_render(self) -> None:
        self.docx.unlink(missing_ok=True)
        (self.run_root / "delivery_assets" / "render_manifest.json").unlink(
            missing_ok=True
        )

    def _fake_canonical_render(
        self,
        combined_path: Path,
        trading_date: str,
        slot: str,
        *,
        root: Path,
        vnext_run_root: Path,
        _vnext_lock_acquired: bool,
    ) -> dict:
        report_record = self.store.load_checkpoint(self.run_root)["completed"]["report"][
            "outputs"
        ]["combined"]
        self.assertEqual(
            Path(combined_path).resolve(),
            (self.root / report_record["path"]).resolve(),
        )
        self.assertEqual((trading_date, slot), ("2026-07-14", "midday"))
        self.assertEqual(Path(root).resolve(), self.root)
        self.assertEqual(Path(vnext_run_root).resolve(), self.run_root)
        self.assertIs(_vnext_lock_acquired, True)
        controller_lock = self.run_root / ".controller.lock"
        self.assertTrue(controller_lock.is_file())
        self.docx.parent.mkdir(parents=True, exist_ok=True)
        self.docx.write_bytes(b"PK-auto-rendered-docx")
        previous = (
            json.loads(
                (
                    self.run_root / "delivery_assets" / "render_manifest.json"
                ).read_text(encoding="utf-8")
            )
            if (
                self.run_root / "delivery_assets" / "render_manifest.json"
            ).is_file()
            else {}
        )
        payload = {
            "schema_version": "1.0",
            "status": "success",
            "trading_date": trading_date,
            "slot": slot,
            "attempts": int(previous.get("attempts", 0)) + 1,
            "source_path": str(Path(combined_path).resolve()),
            "source_sha256": report_record["sha256"],
            "renderer_sha256": canonical_report_renderer.renderer_sha256(),
            "presentation_contract_sha256": sha256_file(
                self.root
                / "governance"
                / "vnext_report_presentation_contract.json"
            ),
            "docx_path": str(self.docx),
            "docx_sha256": sha256_file(self.docx),
        }
        atomic_json_write(
            self.run_root / "delivery_assets" / "render_manifest.json",
            payload,
        )
        return payload

    @staticmethod
    def _fake_render(pdf: Path, prefix: Path, renderer: Path) -> list[Path]:
        assert pdf.is_file()
        assert renderer.name == "pdftoppm.exe"
        pages = []
        for index in (1, 2):
            page = prefix.parent / f"{prefix.name}-{index}.png"
            page.write_bytes(b"\x89PNG\r\n\x1a\n" + bytes([index]))
            pages.append(page)
        return pages

    def test_prepare_and_approve_are_run_local_and_hash_bound(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepared = prepare_run_qa(self.root, self.run_root)

        self.assertEqual(prepared["schema_version"], "word-qa-vnext-v2")
        self.assertEqual(prepared["status"], "prepared")
        self.assertEqual(prepared["backend"], "word_com")
        self.assertEqual(prepared["inspected_pages"], [])
        self.assertEqual(prepared["page_count"], 2)
        self.assertEqual(
            prepared["review_checklist_version"],
            WORD_QA_REVIEW_CHECKLIST_VERSION,
        )
        self.assertEqual(
            prepared["review_checklist"],
            {key: False for key in WORD_QA_REVIEW_CHECKLIST_KEYS},
        )
        self.assertEqual(
            prepared["run_sha256"],
            self.store.load_run(self.run_root)["run_sha256"],
        )
        for record in [prepared["docx"], prepared["pdf"], *prepared["page_images"]]:
            self.assertTrue(record["path"].startswith("runtime_state_vnext/"))
            self.assertNotIn("runtime_state/", record["path"])

        approved = approve_run_qa(
            self.root,
            self.run_root,
            reviewer="Codex",
            confirm_visual_pass=True,
        )
        self.assertEqual(approved["status"], "passed")
        self.assertTrue(approved["visual_qa_pass"])
        self.assertEqual(approved["inspected_pages"], [1, 2])
        self.assertEqual(
            approved["review_checklist"],
            {key: True for key in WORD_QA_REVIEW_CHECKLIST_KEYS},
        )
        self.assertEqual(
            approved["report_sha256"],
            self.store.load_checkpoint(self.run_root)["completed"]["report"]["outputs"]["combined"]["sha256"],
        )

        tampered = dict(approved)
        tampered["review_checklist"] = dict(approved["review_checklist"])
        tampered["review_checklist"]["critical_price_blocks_not_split"] = False
        self.assertFalse(_manifest_files_valid(self.root, tampered))

    def test_word_qa_export_owns_missing_run_local_render(self) -> None:
        self._remove_seeded_render()
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical",
            side_effect=self._fake_canonical_render,
        ) as render_mock, patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepared = prepare_run_qa(self.root, self.run_root)

        render_mock.assert_called_once()
        self.assertEqual(prepared["status"], "prepared")
        self.assertTrue(self.docx.is_file())
        attempts = list(
            (self.run_root / "delivery_assets" / "word_qa").glob("attempt-*")
        )
        self.assertEqual(len(attempts), 1)

    def test_render_failure_is_coded_and_stops_before_word_qa_attempt(self) -> None:
        self._remove_seeded_render()
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical",
            side_effect=RuntimeError("synthetic render failure"),
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word"
        ) as word_export:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)

        failure = captured.exception
        self.assertEqual(failure.failure_class, FailureClass.DELIVERY_INTEGRITY)
        self.assertEqual(
            [defect["code"] for defect in failure.defects],
            ["delivery.render_failed_before_word_qa"],
        )
        word_export.assert_not_called()
        self.assertFalse(
            (self.run_root / "delivery_assets" / "word_qa.json").exists()
        )
        word_qa_root = self.run_root / "delivery_assets" / "word_qa"
        self.assertEqual(
            list(word_qa_root.glob("attempt-*")) if word_qa_root.exists() else [],
            [],
        )
        events = [
            json.loads(line)
            for line in (self.run_root / "events.jsonl").read_text(
                encoding="utf-8"
            ).splitlines()
        ]
        self.assertEqual(events[-1]["event"], "delivery_preparation_failed")
        self.assertEqual(
            events[-1]["failure_code"],
            "delivery.render_failed_before_word_qa",
        )

    def test_malformed_failed_render_is_nonretryable_contract_damage(self) -> None:
        self._remove_seeded_render()

        def malformed_failure(*args, **kwargs) -> None:
            atomic_json_write(
                self.run_root / "delivery_assets" / "render_manifest.json",
                {
                    "schema_version": "1.0",
                    "status": "failed",
                    "attempts": 1,
                },
            )
            raise RuntimeError("synthetic malformed render failure")

        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical",
            side_effect=malformed_failure,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word"
        ) as word_export:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)

        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.MANUAL_REVIEW,
        )
        self.assertFalse(captured.exception.policy.retryable)
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.render_contract_invalid"],
        )
        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(state["state_code"], "delivery.render_contract_invalid")
        self.assertEqual(
            state["next_action_code"],
            "delivery.maintenance_required",
        )
        self.assertIsNone(state["next_command"])
        word_export.assert_not_called()

    def test_delivery_control_state_owns_the_next_action(self) -> None:
        pending = delivery_control_state(self.root, self.run_root)
        self.assertEqual(pending["state_code"], "delivery.word_qa_export_pending")
        self.assertEqual(pending["next_action_code"], "delivery.word_qa_export")
        self.assertEqual(pending["next_command"], "word-qa-export")

        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepare_run_qa(self.root, self.run_root)

        review = delivery_control_state(self.root, self.run_root)
        self.assertEqual(review["state_code"], "delivery.word_qa_review_pending")
        self.assertEqual(
            review["next_action_code"], "delivery.word_qa_visual_review"
        )
        self.assertIsNone(review["next_command"])

        approve_run_qa(
            self.root,
            self.run_root,
            reviewer="Codex",
            confirm_visual_pass=True,
        )
        ready = delivery_control_state(self.root, self.run_root)
        self.assertEqual(
            ready["state_code"], "delivery.shadow_delivery_pending"
        )
        self.assertEqual(
            ready["next_action_code"], "delivery.shadow_delivery"
        )
        self.assertEqual(ready["next_command"], "shadow-delivery")

    def test_delivery_control_state_rejects_invalid_render_binding(self) -> None:
        manifest_path = self.run_root / "delivery_assets" / "render_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["source_sha256"] = "0" * 64
        atomic_json_write(manifest_path, manifest)

        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(state["state_code"], "delivery.render_contract_invalid")
        self.assertEqual(
            state["next_action_code"],
            "delivery.maintenance_required",
        )
        self.assertIsNone(state["next_command"])
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical"
        ) as render_mock:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)
        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.MANUAL_REVIEW,
        )
        self.assertFalse(captured.exception.policy.retryable)
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.render_contract_invalid"],
        )
        render_mock.assert_not_called()

    def test_delivery_control_state_rejects_cross_run_passed_word_qa(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepare_run_qa(self.root, self.run_root)
        approve_run_qa(
            self.root,
            self.run_root,
            reviewer="Codex",
            confirm_visual_pass=True,
        )
        manifest_path = self.run_root / "delivery_assets" / "word_qa.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["run_id"] = "run-another-owner"
        atomic_json_write(manifest_path, manifest)

        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(
            state["state_code"],
            "delivery.word_qa_contract_invalid",
        )
        self.assertEqual(
            state["next_action_code"],
            "delivery.maintenance_required",
        )
        self.assertIsNone(state["next_command"])

    def test_approve_and_reject_refuse_invalid_prepared_contract(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepare_run_qa(self.root, self.run_root)

        manifest_path = self.run_root / "delivery_assets" / "word_qa.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["attempts"] = "1"
        atomic_json_write(manifest_path, manifest)
        invalid_bytes = manifest_path.read_bytes()

        for operation in (
            lambda: approve_run_qa(
                self.root,
                self.run_root,
                reviewer="Codex",
                confirm_visual_pass=True,
            ),
            lambda: reject_run_qa(
                self.root,
                self.run_root,
                reason="invalid prepared contract must not be promoted",
            ),
        ):
            with self.subTest(operation=operation):
                with self.assertRaises(WorkflowFailure) as captured:
                    operation()
                self.assertEqual(
                    captured.exception.failure_class,
                    FailureClass.MANUAL_REVIEW,
                )
                self.assertFalse(captured.exception.policy.retryable)
                self.assertEqual(
                    [item["code"] for item in captured.exception.defects],
                    ["delivery.word_qa_contract_invalid"],
                )
                self.assertEqual(manifest_path.read_bytes(), invalid_bytes)

    def test_delivery_control_state_rejects_cross_run_page_evidence(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepare_run_qa(self.root, self.run_root)

        manifest_path = self.run_root / "delivery_assets" / "word_qa.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        foreign_attempt = (
            self.root
            / "runtime_state_vnext"
            / "foreign-run"
            / "delivery_assets"
            / "word_qa"
            / "attempt-001-deadbeef"
        )
        foreign_attempt.mkdir(parents=True)
        source_pdf = self.root / manifest["pdf"]["path"]
        foreign_pdf = foreign_attempt / source_pdf.name
        foreign_pdf.write_bytes(source_pdf.read_bytes())
        manifest["pdf"] = _artifact_record(self.root, foreign_pdf)
        manifest["pdf_sha256"] = manifest["pdf"]["sha256"]
        foreign_pages = []
        for page in manifest["page_images"]:
            source_page = self.root / page["path"]
            foreign_page = foreign_attempt / source_page.name
            foreign_page.write_bytes(source_page.read_bytes())
            record = _artifact_record(self.root, foreign_page)
            record["page"] = page["page"]
            foreign_pages.append(record)
        manifest["page_images"] = foreign_pages
        atomic_json_write(manifest_path, manifest)

        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(
            state["state_code"],
            "delivery.word_qa_contract_invalid",
        )
        self.assertEqual(
            state["next_action_code"],
            "delivery.maintenance_required",
        )
        self.assertIsNone(state["next_command"])

    def test_cli_emits_exact_render_failure_code(self) -> None:
        from runtime_vnext import cli

        self._remove_seeded_render()
        stdout = io.StringIO()
        argv = [
            "runtime_vnext.cli",
            "word-qa-export",
            "--date",
            "2026-07-14",
            "--slot",
            "midday",
            "--run-id",
            self.run_id,
        ]
        with patch.object(cli, "ROOT", self.root), patch(
            "sys.argv", argv
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical",
            side_effect=RuntimeError("synthetic render failure"),
        ), redirect_stdout(stdout):
            exit_code = cli.main()

        payload = json.loads(stdout.getvalue())
        self.assertEqual(exit_code, 2)
        self.assertEqual(payload["failure_class"], "delivery_integrity")
        self.assertEqual(
            payload["failure_code"],
            "delivery.render_failed_before_word_qa",
        )
        self.assertEqual(
            payload["defect_codes"],
            ["delivery.render_failed_before_word_qa"],
        )

    def test_stale_renderer_identity_rerenders_after_visual_rejection(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            first = prepare_run_qa(self.root, self.run_root)
        reject_run_qa(
            self.root,
            self.run_root,
            reason="visual layout requires a renderer revision",
        )

        render_manifest_path = (
            self.run_root / "delivery_assets" / "render_manifest.json"
        )
        render_manifest = json.loads(
            render_manifest_path.read_text(encoding="utf-8")
        )
        render_manifest["renderer_sha256"] = "0" * 64
        atomic_json_write(render_manifest_path, render_manifest)
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical",
            side_effect=self._fake_canonical_render,
        ) as render_mock, patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            second = prepare_run_qa(self.root, self.run_root)

        render_mock.assert_called_once()
        self.assertEqual(first["attempts"], 1)
        self.assertEqual(second["attempts"], 2)
        self.assertNotEqual(first["docx_sha256"], second["docx_sha256"])

    def test_malformed_renderer_identity_requires_maintenance(self) -> None:
        manifest_path = self.run_root / "delivery_assets" / "render_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["renderer_sha256"] = "not-a-sha256"
        atomic_json_write(manifest_path, manifest)

        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(state["state_code"], "delivery.render_contract_invalid")
        self.assertEqual(
            state["next_action_code"],
            "delivery.maintenance_required",
        )
        self.assertIsNone(state["next_command"])
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical"
        ) as render_mock:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)
        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.MANUAL_REVIEW,
        )
        self.assertFalse(captured.exception.policy.retryable)
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.render_contract_invalid"],
        )
        render_mock.assert_not_called()

    def test_rejected_docx_hash_tamper_cannot_fake_a_render_revision(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepare_run_qa(self.root, self.run_root)
        reject_run_qa(
            self.root,
            self.run_root,
            reason="visual layout requires a real renderer revision",
        )
        manifest_path = self.run_root / "delivery_assets" / "word_qa.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["docx"]["sha256"] = "0" * 64
        manifest["docx_sha256"] = "0" * 64
        atomic_json_write(manifest_path, manifest)

        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(
            state["state_code"],
            "delivery.word_qa_contract_invalid",
        )
        self.assertEqual(
            state["next_action_code"],
            "delivery.maintenance_required",
        )
        self.assertIsNone(state["next_command"])
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical"
        ) as render_mock, patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word"
        ) as word_export:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)
        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.MANUAL_REVIEW,
        )
        self.assertFalse(captured.exception.policy.retryable)
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.word_qa_contract_invalid"],
        )
        render_mock.assert_not_called()
        word_export.assert_not_called()

    def test_render_attempt_limit_has_no_retry_command(self) -> None:
        self.docx.unlink()
        manifest_path = self.run_root / "delivery_assets" / "render_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest.update(
            {
                "status": "failed",
                "attempts": 2,
                "failure_reason": "synthetic exhausted render",
            }
        )
        manifest.pop("docx_sha256", None)
        atomic_json_write(manifest_path, manifest)

        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical"
        ) as render_mock, patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word"
        ) as word_export:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)

        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.MANUAL_REVIEW,
        )
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.render_attempt_limit_reached"],
        )
        render_mock.assert_not_called()
        word_export.assert_not_called()
        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(
            state["state_code"],
            "delivery.render_attempt_limit_reached",
        )
        self.assertIsNone(state["next_command"])

    def test_word_qa_attempt_limit_has_no_retry_command(self) -> None:
        run = self.store.load_run(self.run_root)
        report = self.store.load_checkpoint(self.run_root)["completed"]["report"][
            "outputs"
        ]["combined"]
        atomic_json_write(
            self.run_root / "delivery_assets" / "word_qa.json",
            {
                "schema_version": "word-qa-vnext-v2",
                "status": "failed",
                "visual_qa_pass": False,
                "created_at": "2026-07-14T12:00:00+08:00",
                "trading_date": run["trading_date"],
                "slot": run["slot"],
                "run_id": run["run_id"],
                "run_sha256": run["run_sha256"],
                "attempts": 2,
                "backend": "word_com",
                "report_sha256": report["sha256"],
                "docx": _artifact_record(self.root, self.docx),
                "docx_sha256": sha256_file(self.docx),
                "failure_reason": "synthetic exhausted Word export",
            },
        )
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word"
        ) as word_export:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)

        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.MANUAL_REVIEW,
        )
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.word_qa_attempt_limit_reached"],
        )
        word_export.assert_not_called()
        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(
            state["state_code"],
            "delivery.word_qa_attempt_limit_reached",
        )
        self.assertIsNone(state["next_command"])

    def test_first_word_qa_failure_is_coded_and_retryable(self) -> None:
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=RuntimeError("synthetic Word export failure"),
        ):
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)

        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.DELIVERY_INTEGRITY,
        )
        self.assertTrue(captured.exception.policy.retryable)
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.word_qa_export_failed"],
        )
        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(state["state_code"], "delivery.word_qa_export_pending")
        self.assertEqual(state["next_command"], "word-qa-export")

    def test_failed_word_export_hash_tamper_requires_maintenance(self) -> None:
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=RuntimeError("synthetic Word export failure"),
        ):
            with self.assertRaises(WorkflowFailure):
                prepare_run_qa(self.root, self.run_root)
        manifest_path = self.run_root / "delivery_assets" / "word_qa.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["docx"]["sha256"] = "0" * 64
        manifest["docx_sha256"] = "0" * 64
        atomic_json_write(manifest_path, manifest)

        state = delivery_control_state(self.root, self.run_root)
        self.assertEqual(
            state["state_code"],
            "delivery.word_qa_contract_invalid",
        )
        self.assertEqual(
            state["next_action_code"],
            "delivery.maintenance_required",
        )
        self.assertIsNone(state["next_command"])
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical"
        ) as render_mock, patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word"
        ) as word_export:
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)
        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.MANUAL_REVIEW,
        )
        self.assertFalse(captured.exception.policy.retryable)
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.word_qa_contract_invalid"],
        )
        render_mock.assert_not_called()
        word_export.assert_not_called()

    def test_malformed_page_output_never_becomes_prepared_qa(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")

        def malformed_pages(
            pdf: Path,
            prefix: Path,
            page_renderer: Path,
        ) -> list[Path]:
            self.assertTrue(pdf.is_file())
            self.assertEqual(page_renderer, renderer)
            pages = []
            for page_number in (1, 3):
                page = prefix.parent / f"{prefix.name}-{page_number}.png"
                page.write_bytes(b"\x89PNG\r\n\x1a\n" + bytes([page_number]))
                pages.append(page)
            return pages

        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=malformed_pages,
        ):
            with self.assertRaises(WorkflowFailure) as captured:
                prepare_run_qa(self.root, self.run_root)

        self.assertEqual(
            captured.exception.failure_class,
            FailureClass.DELIVERY_INTEGRITY,
        )
        self.assertTrue(captured.exception.policy.retryable)
        self.assertEqual(
            [item["code"] for item in captured.exception.defects],
            ["delivery.word_qa_export_failed"],
        )
        manifest = json.loads(
            (
                self.run_root / "delivery_assets" / "word_qa.json"
            ).read_text(encoding="utf-8")
        )
        self.assertEqual(manifest["status"], "failed")
        self.assertEqual(manifest["attempts"], 1)
        self.assertNotIn("page_images", manifest)

    def test_certified_report_tamper_stops_before_render_and_word(self) -> None:
        self._remove_seeded_render()
        report_record = self.store.load_checkpoint(self.run_root)["completed"][
            "report"
        ]["outputs"]["combined"]
        (self.root / report_record["path"]).write_text(
            "tampered report",
            encoding="utf-8",
        )
        with patch(
            "runtime_vnext.delivery_renderer.canonical_report_renderer.render_canonical"
        ) as render_mock, patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word"
        ) as word_export:
            with self.assertRaisesRegex(ValueError, "certified report"):
                prepare_run_qa(self.root, self.run_root)

        render_mock.assert_not_called()
        word_export.assert_not_called()
        self.assertFalse(
            (self.run_root / "delivery_assets" / "word_qa").exists()
        )

    def test_docx_semantic_pagination_marks_sources_and_price_blocks(self) -> None:
        markdown = self.root / "semantic-pagination.md"
        output = self.root / "semantic-pagination.docx"
        markdown.write_text(
            "\n".join(
                [
                    "# 语义分页测试",
                    "**1. BlockBeats｜重要事件事实与影响层级**",
                    "证据事实：事件事实。",
                    "策略：入场区间63960–64120分档做空。",
                    "止盈：TP1为63000；TP2为62000。",
                    "止损：成交后止损为64500。",
                    "第一层：触及63200后执行防守。",
                    "第二层：触及62700后移至保本。",
                    "注意事项：成交后按实际加权成交价重算。",
                    "取消价≠止损：未成交前触及64600只取消订单。",
                ]
            ),
            encoding="utf-8",
        )
        render(markdown, output)
        document = Document(output)
        paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
        source = paragraphs["1. BlockBeats｜重要事件事实与影响层级"]
        self.assertTrue(source.paragraph_format.keep_with_next)
        self.assertTrue(source.paragraph_format.keep_together)
        critical = {}
        for prefix in (
            "策略：",
            "止盈：",
            "止损：",
            "第一层：",
            "第二层：",
            "注意事项：",
            "取消价≠止损：",
        ):
            paragraph = next(
                item for text, item in paragraphs.items() if text.startswith(prefix)
            )
            critical[prefix] = paragraph
            with self.subTest(prefix=prefix):
                self.assertTrue(paragraph.paragraph_format.keep_together)
        for prefix in ("策略：", "止盈：", "第一层：", "第二层：", "注意事项："):
            with self.subTest(keep_with_next=prefix):
                self.assertTrue(critical[prefix].paragraph_format.keep_with_next)
        for prefix in ("止损：", "取消价≠止损："):
            with self.subTest(chain_tail=prefix):
                self.assertIsNot(critical[prefix].paragraph_format.keep_with_next, True)

    def test_docx_renderer_preserves_critical_price_text_exactly(self) -> None:
        markdown = self.root / "arabic-price-preservation.md"
        output = self.root / "arabic-price-preservation.docx"
        expected = [
            "策略：候选入场区间63960–64120分档做多；候选挂单参数为63960、64120。",
            "止盈：候选TP1为65050；候选TP2为65500。",
            "止损：候选初始止损为63750。",
            "第一层：候选触发价为63650。",
            "第二层：候选触发价为63520。",
            "注意事项：现价63960与区间64120均保持阿拉伯数字。",
            "取消价≠止损：候选取消价为64280。",
        ]
        markdown.write_text("\n".join(["# 阿拉伯数字保真测试", *expected]), encoding="utf-8")

        render(markdown, output)
        rendered_text = [paragraph.text for paragraph in Document(output).paragraphs]
        for line in expected:
            with self.subTest(line=line):
                self.assertIn(line, rendered_text)
        combined = "\n".join(rendered_text)
        for number in (
            "63960",
            "64120",
            "65050",
            "65500",
            "63750",
            "63650",
            "63520",
            "64280",
        ):
            with self.subTest(number=number):
                self.assertIn(number, combined)

    def test_rejected_candidate_audit_uses_neutral_non_trade_styling(self) -> None:
        markdown = self.root / "rejected-candidate-audit.md"
        output = self.root / "rejected-candidate-audit.docx"
        risk_scenarios = [
            (
                "被拒候选风险情景（不可执行，不得用于下单）："
                f"成交档位{index}；加权入场66240；名义风险13.29%。"
            )
            for index in range(1, 32)
        ]
        expected = [
            "被拒候选否决原因（不可执行，不得用于下单）：1.结构止损无效；2.风险数学无效。",
            "被拒候选审计输入（不可执行，不得用于下单）：低多；候选名义持仓=750 USDT；止损=64480。",
            *risk_scenarios,
            "被拒候选机械结论（不可执行，不得用于下单）：不得创建真实订单。",
        ]
        markdown.write_text(
            "\n".join(["# 被拒候选中性样式测试", *expected]),
            encoding="utf-8",
        )

        render(markdown, output)
        paragraphs = [
            paragraph
            for paragraph in Document(output).paragraphs
            if paragraph.text.startswith("被拒候选")
            and "（不可执行，不得用于下单）" in paragraph.text
        ]
        self.assertEqual([paragraph.text for paragraph in paragraphs], expected)
        for index, paragraph in enumerate(paragraphs):
            with self.subTest(index=index):
                colors = {
                    str(run.font.color.rgb)
                    for run in paragraph.runs
                    if run.text and run.font.color.rgb is not None
                }
                self.assertEqual(colors, {"666666"})
                self.assertTrue(paragraph.paragraph_format.keep_together)
                if index == 0:
                    self.assertTrue(paragraph.paragraph_format.keep_with_next)
                else:
                    self.assertIsNot(
                        paragraph.paragraph_format.keep_with_next,
                        True,
                    )
        combined = "\n".join(paragraph.text for paragraph in paragraphs)
        for number in ("750", "64480", "66240", "13.29"):
            with self.subTest(number=number):
                self.assertIn(number, combined)

    def test_critical_price_keep_chains_do_not_bridge_unrelated_paragraphs(self) -> None:
        markdown = self.root / "broken-price-chain.md"
        output = self.root / "broken-price-chain.docx"
        markdown.write_text(
            "\n".join(
                [
                    "# 断链测试",
                    "策略：入场区间63960–64120分档做空。",
                    "普通说明：此段故意打断连续价格组。",
                    "止盈：TP1为63000；TP2为62000。",
                    "止损：成交后止损为64500。",
                ]
            ),
            encoding="utf-8",
        )

        render(markdown, output)
        strategy = next(
            paragraph
            for paragraph in Document(output).paragraphs
            if paragraph.text.startswith("策略：")
        )
        self.assertIsNot(strategy.paragraph_format.keep_with_next, True)

    def test_visual_rejection_blocks_unchanged_docx(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepare_run_qa(self.root, self.run_root)
        reject_run_qa(
            self.root,
            self.run_root,
            reason="page layout requires correction",
        )
        with self.assertRaisesRegex(RuntimeError, "unchanged"):
            prepare_run_qa(self.root, self.run_root)

    def test_python_approve_requires_explicit_visual_confirmation(self) -> None:
        renderer = self.root / "pdftoppm.exe"
        renderer.write_bytes(b"binary")
        with patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._export_word",
            side_effect=self._fake_export,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._find_pdftoppm",
            return_value=renderer,
        ), patch(
            "runtime_vnext.delivery_renderer.canonical_word_qa._render_pdf",
            side_effect=self._fake_render,
        ):
            prepare_run_qa(self.root, self.run_root)
        manifest_path = self.run_root / "delivery_assets" / "word_qa.json"
        before = manifest_path.read_bytes()

        with self.assertRaisesRegex(ValueError, "confirm_visual_pass=True"):
            approve_run_qa(self.root, self.run_root, reviewer="Codex")

        self.assertEqual(manifest_path.read_bytes(), before)
        self.assertEqual(
            json.loads(manifest_path.read_text(encoding="utf-8"))["status"],
            "prepared",
        )

    def test_legacy_nonterminal_run_is_read_only_for_all_word_qa_mutations(self) -> None:
        run_path = self.run_root / "run.json"
        run = json.loads(run_path.read_text(encoding="utf-8"))
        run["schema_version"] = "daily-run-vnext-v1"
        run.pop("run_mode")
        run.pop("authorization_kind")
        run.pop("authorization_ref_sha256")
        run.pop("run_sha256")
        run["run_sha256"] = canonical_sha256(run)
        atomic_json_write(run_path, run)
        before = {
            path.relative_to(self.run_root).as_posix(): path.read_bytes()
            for path in self.run_root.rglob("*")
            if path.is_file()
        }

        operations = (
            lambda: prepare_run_qa(self.root, self.run_root),
            lambda: approve_run_qa(
                self.root,
                self.run_root,
                reviewer="Codex",
                confirm_visual_pass=True,
            ),
            lambda: reject_run_qa(
                self.root,
                self.run_root,
                reason="legacy run must remain immutable",
            ),
        )
        for operation in operations:
            with self.subTest(operation=operation):
                with self.assertRaisesRegex(ValueError, "read-only"):
                    operation()

        after = {
            path.relative_to(self.run_root).as_posix(): path.read_bytes()
            for path in self.run_root.rglob("*")
            if path.is_file()
        }
        self.assertEqual(after, before)

    def test_word_qa_mutations_require_the_same_active_run_identity(self) -> None:
        active_path = self.root / "runtime_state_vnext" / "active_run.json"
        original_active = json.loads(active_path.read_text(encoding="utf-8"))
        before = {
            path.relative_to(self.run_root).as_posix(): path.read_bytes()
            for path in self.run_root.rglob("*")
            if path.is_file()
        }
        operations = (
            lambda: prepare_run_qa(self.root, self.run_root),
            lambda: approve_run_qa(
                self.root,
                self.run_root,
                reviewer="Codex",
                confirm_visual_pass=True,
            ),
            lambda: reject_run_qa(
                self.root,
                self.run_root,
                reason="inactive run must remain immutable",
            ),
        )

        active_path.unlink()
        for operation in operations:
            with self.subTest(active_state="missing", operation=operation):
                with self.assertRaisesRegex(RuntimeError, "active-run lease is missing"):
                    operation()

        wrong_active = dict(original_active)
        wrong_active["run_id"] = "run-another-owner"
        wrong_active.pop("active_sha256")
        wrong_active["active_sha256"] = canonical_sha256(wrong_active)
        atomic_json_write(active_path, wrong_active)
        for operation in operations:
            with self.subTest(active_state="wrong_owner", operation=operation):
                with self.assertRaisesRegex(RuntimeError, "belongs to another run"):
                    operation()

        after = {
            path.relative_to(self.run_root).as_posix(): path.read_bytes()
            for path in self.run_root.rglob("*")
            if path.is_file()
        }
        self.assertEqual(after, before)

    def test_producer_source_has_no_alternate_office_backend(self) -> None:
        source = (
            Path(__file__).resolve().parents[1]
            / "runtime_vnext"
            / "delivery_renderer.py"
        ).read_text(encoding="utf-8").casefold()
        self.assertNotIn("soffice", source)
        self.assertNotIn("libreoffice", source)


if __name__ == "__main__":
    unittest.main()
